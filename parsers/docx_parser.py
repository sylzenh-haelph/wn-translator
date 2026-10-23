from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn

from models.document import Document, Paragraph, TextRun


def _run_formatting(run):
    return {
        "bold": bool(run.bold),
        "italic": bool(run.italic),
        "underline": bool(run.underline),
        "font_name": run.font.name,
        "font_size": (
            run.font.size.pt
            if run.font.size is not None
            else None
        ),
        "font_color": (
            str(run.font.color.rgb)
            if run.font.color.rgb is not None
            else None
        ),
        "superscript": bool(run.font.superscript),
        "subscript": bool(run.font.subscript),
    }


def _paragraph_pictures(source_paragraph, source):
    pictures = []
    for blip in source_paragraph._p.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid or rid not in source.part.rels:
            continue
        part = source.part.rels[rid].target_part
        suffix = Path(str(part.partname)).suffix.lower() or ".png"
        pictures.append((part.blob, suffix))
    return pictures


def parse_docx(path):
    path = Path(path)

    source = DocxDocument(path)

    document = Document()

    document.title = path.stem

    images = {}

    paragraph_number = 0

    for source_paragraph in source.paragraphs:
        # Jangan membuang paragraf kosong.
        # Posisi paragraf harus tetap dipertahankan.
        paragraph_id = f"p{paragraph_number:04d}"
        paragraph_number += 1

        runs = []

        for source_run in source_paragraph.runs:
            runs.append(
                TextRun(
                    text=source_run.text,
                    formatting=_run_formatting(source_run),
                )
            )

        style = {
            "name": (
                source_paragraph.style.name
                if source_paragraph.style is not None
                else None
            ),
            "alignment": (
                str(source_paragraph.alignment)
                if source_paragraph.alignment is not None
                else None
            ),
        }

        image_names = []
        for blob, suffix in _paragraph_pictures(source_paragraph, source):
            name = f"{paragraph_id}_{len(image_names)}{suffix}"
            images[name] = blob
            image_names.append(name)
        if image_names:
            style["images"] = image_names

        document.paragraphs.append(
            Paragraph(
                id=paragraph_id,
                runs=runs,
                style=style,
            )
        )

    document.metadata["docx_images"] = images

    return document
