from pathlib import Path
import io
import struct
import sys
import zlib

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document as WordDocument
from docx.shared import Inches

from models.document import Paragraph
from parsers.docx_parser import parse_docx
from reconstruction.epub_reconstructor import _paragraph_html


def make_png():
    def chunk(tag, data):
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = b"\x00\xff\x00\x00"
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def test_docx_picture_is_extracted(tmp_path):
    path = tmp_path / "a.docx"
    word = WordDocument()
    word.add_paragraph("Teks")
    word.add_picture(io.BytesIO(make_png()), width=Inches(1))
    word.save(path)

    parsed = parse_docx(path)
    images = parsed.metadata.get("docx_images", {})
    assert len(images) == 1
    name = next(iter(images))
    assert name.endswith(".png")
    assert any(p.style.get("images") == [name] for p in parsed.paragraphs)


def test_image_paragraph_renders_centered_img():
    paragraph = Paragraph(id="p1", runs=[], style={"images": ["p1_0.png"]})
    html = _paragraph_html(paragraph)
    assert 'src="images/p1_0.png"' in html
    assert "margin:0 auto" in html
