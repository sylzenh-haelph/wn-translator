import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor

from parsers.docx_parser import parse_docx
from reconstruction.docx_reconstructor import reconstruct_docx


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "test_docx_full_fidelity_source.docx"
OUTPUT = ROOT / "test_docx_full_fidelity_output.docx"


def create_fixture():
    document = DocxDocument()

    heading = document.add_paragraph()
    heading.style = "Heading 1"

    run = heading.add_run("Chapter 1")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(18)

    paragraph = document.add_paragraph()

    run = paragraph.add_run("Alice ")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(12)

    run = paragraph.add_run("enters")
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    run = paragraph.add_run(" the palace.")
    run.underline = True
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x00, 0x66, 0x99)

    paragraph = document.add_paragraph()

    run = paragraph.add_run("H")
    run.superscript = True

    run = paragraph.add_run("2")
    run.subscript = True

    document.save(SOURCE)


def main():
    # ============================================================
    # 1. Create controlled source DOCX
    # ============================================================

    create_fixture()
    assert SOURCE.exists()

    # ============================================================
    # 2. Parse source
    # ============================================================

    source = parse_docx(SOURCE)

    assert len(source.paragraphs) == 3

    source_texts = [
        paragraph.text
        for paragraph in source.paragraphs
    ]

    assert source_texts == [
        "Chapter 1",
        "Alice enters the palace.",
        "H2",
    ]

    # ============================================================
    # 3. Reconstruct using the parsed document
    # ============================================================

    reconstruct_docx(
        source,
        OUTPUT,
    )

    assert OUTPUT.exists()
    assert OUTPUT.stat().st_size > 0

    # ============================================================
    # 4. Parse reconstructed DOCX
    # ============================================================

    result = parse_docx(OUTPUT)

    assert len(result.paragraphs) == 3

    result_texts = [
        paragraph.text
        for paragraph in result.paragraphs
    ]

    assert result_texts == source_texts

    # ============================================================
    # 5. Verify paragraph styles
    # ============================================================

    assert (
        result.paragraphs[0].style.get("name")
        == source.paragraphs[0].style.get("name")
    )

    # ============================================================
    # 6. Verify run formatting
    # ============================================================

    source_runs = [
        run
        for paragraph in source.paragraphs
        for run in paragraph.runs
    ]

    result_runs = [
        run
        for paragraph in result.paragraphs
        for run in paragraph.runs
    ]

    assert len(result_runs) == len(source_runs)

    for source_run, result_run in zip(
        source_runs,
        result_runs,
    ):
        assert result_run.text == source_run.text

        for key in [
            "bold",
            "italic",
            "underline",
            "font_name",
            "font_size",
            "font_color",
            "superscript",
            "subscript",
        ]:
            assert (
                result_run.formatting.get(key)
                == source_run.formatting.get(key)
            ), (
                f"Formatting mismatch for {key}: "
                f"{source_run.formatting.get(key)!r} != "
                f"{result_run.formatting.get(key)!r}"
            )

    # ============================================================
    # 7. Verify physical DOCX can be opened by python-docx
    # ============================================================

    reopened = DocxDocument(OUTPUT)

    assert len(reopened.paragraphs) == 3
    assert reopened.paragraphs[0].text == "Chapter 1"
    assert reopened.paragraphs[1].text == "Alice enters the palace."
    assert reopened.paragraphs[2].text == "H2"

    # ============================================================
    # 8. Cleanup
    # ============================================================

    SOURCE.unlink(missing_ok=True)
    OUTPUT.unlink(missing_ok=True)

    print("PASS: DOCX full fidelity regression")


if __name__ == "__main__":
    main()
