from pathlib import Path
import io
import struct
import sys
import xml.etree.ElementTree as ET
import zlib
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document as WordDocument
from docx.shared import Inches

from models.document import TextRun
from parsers.docx_parser import parse_docx
from reconstruction.epub_reconstructor import reconstruct_epub
from reconstruction.epub_validator import validate_epub

CRITICAL = {"XHTML_XML", "OPF_RESOURCE_EXISTS", "NAV_FILE", "NAV_EXISTS"}


def make_png():
    def chunk(tag, data):
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
        + chunk(b"IEND", b"")
    )


def build_sample_docx(path):
    word = WordDocument()
    word.add_heading("Chapter 1: Ujian", level=1)
    word.add_paragraph("Subaru & Al: [Halo, dunia.]")
    word.add_paragraph("")
    word.add_paragraph('"Ayo pergi," kata Emilia.')
    word.add_picture(io.BytesIO(make_png()), width=Inches(1))
    word.add_paragraph("Akhir bab.")
    word.save(path)


def fake_translate(document):
    for paragraph in document.paragraphs:
        text = "".join(run.text for run in paragraph.runs)
        if text.strip():
            paragraph.runs = [TextRun(text="[ID] " + text)]
    return document


def test_docx_to_epub_end_to_end(tmp_path):
    src = tmp_path / "bab.docx"
    build_sample_docx(src)

    translated = fake_translate(parse_docx(src))

    out = tmp_path / "bab_id.epub"
    reconstruct_epub(translated, out, source_path=None)

    result = validate_epub(out)
    rules = [issue.rule for issue in result.issues]
    assert not [rule for rule in rules if rule in CRITICAL], rules

    with ZipFile(out) as archive:
        names = archive.namelist()
        assert any(name.startswith("OEBPS/images/") and name.endswith(".png") for name in names)
        chapter = next(n for n in names if n.endswith(".xhtml") and "content" in n)
        markup = archive.read(chapter).decode("utf-8")

    root = ET.fromstring(markup.encode("utf-8"))
    text = "".join(root.itertext())
    assert "[ID] Subaru & Al: [Halo, dunia.]" in text
    assert "[ID] Akhir bab." in text
    assert "<img" in markup
