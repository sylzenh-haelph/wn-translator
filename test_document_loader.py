import shutil
from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Pt

from parsers.document_loader import load_document


TEST_DIR = Path("document_loader_test")

if TEST_DIR.exists():
    shutil.rmtree(TEST_DIR)

TEST_DIR.mkdir(parents=True)


# ----------------------------------------
# CREATE TEST DOCX
# ----------------------------------------

docx_path = TEST_DIR / "test_book.docx"

doc = DocxDocument()

# Paragraph 1: heading
heading = doc.add_paragraph()
heading.style = "Heading 1"

run = heading.add_run("Chapter 1")
run.bold = True

# Paragraph 2
paragraph = doc.add_paragraph()

run = paragraph.add_run(
    "Alice enters the palace."
)
run.bold = True
run.font.size = Pt(12)

# Paragraph 3
doc.add_paragraph(
    "Marcus follows her."
)

doc.save(docx_path)


# ----------------------------------------
# LOAD DOCX
# ----------------------------------------

print("=== LOAD DOCX ===")

document = load_document(docx_path)

print("Title:", document.title)
print("Author:", document.author)
print(
    "Paragraph count:",
    len(document.paragraphs),
)

assert document.title == "test_book"
assert len(document.paragraphs) == 3


# ----------------------------------------
# TEXT CHECK
# ----------------------------------------

print("\n=== TEXT ===")

texts = [
    paragraph.text
    for paragraph in document.paragraphs
]

print(texts)

assert texts == [
    "Chapter 1",
    "Alice enters the palace.",
    "Marcus follows her.",
]


# ----------------------------------------
# HEADING CHECK
# ----------------------------------------

print("\n=== HEADING ===")

heading = document.paragraphs[0]

print(
    "Style:",
    heading.style["name"],
)

assert heading.style["name"] == "Heading 1"


# ----------------------------------------
# FORMATTING CHECK
# ----------------------------------------

print("\n=== FORMATTING ===")

paragraph = document.paragraphs[1]

print(
    "Text:",
    paragraph.text,
)

print(
    "Bold:",
    paragraph.runs[0].formatting["bold"],
)

print(
    "Font size:",
    paragraph.runs[0].formatting["font_size"],
)

assert paragraph.runs[0].formatting[
    "bold"
] is True

assert paragraph.runs[0].formatting[
    "font_size"
] == 12


# ----------------------------------------
# NORMAL PARAGRAPH CHECK
# ----------------------------------------

print("\n=== NORMAL PARAGRAPH ===")

normal = document.paragraphs[2]

print(
    "Text:",
    normal.text,
)

assert normal.text == "Marcus follows her."


# ----------------------------------------
# UNSUPPORTED FORMAT
# ----------------------------------------

print("\n=== UNSUPPORTED FORMAT ===")

bad_path = TEST_DIR / "test.txt"

bad_path.write_text(
    "test",
    encoding="utf-8",
)

try:
    load_document(bad_path)
except ValueError as error:
    print("Error:", error)
else:
    raise AssertionError(
        "Format .txt seharusnya ditolak."
    )


print("\nPASS")

shutil.rmtree(TEST_DIR)
