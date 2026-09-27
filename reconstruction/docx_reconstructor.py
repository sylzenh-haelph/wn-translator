from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor

from models.document import Document


def _apply_formatting(target_run, formatting):
    target_run.bold = formatting.get("bold", False)
    target_run.italic = formatting.get("italic", False)
    target_run.underline = formatting.get("underline", False)

    target_run.font.name = formatting.get("font_name")

    font_size = formatting.get("font_size")
    if font_size is not None:
        target_run.font.size = Pt(font_size)

    font_color = formatting.get("font_color")
    if font_color:
        try:
            target_run.font.color.rgb = RGBColor.from_string(font_color)
        except ValueError:
            pass

    target_run.font.superscript = formatting.get(
        "superscript", False
    )
    target_run.font.subscript = formatting.get(
        "subscript", False
    )


def reconstruct_docx(document: Document, output_path):
    output_path = Path(output_path)

    output = DocxDocument()

    for paragraph in document.paragraphs:
        target_paragraph = output.add_paragraph()

        # Pertahankan style paragraf jika memungkinkan.
        style_name = paragraph.style.get("name")

        if style_name:
            try:
                target_paragraph.style = style_name
            except KeyError:
                pass

        for source_run in paragraph.runs:
            target_run = target_paragraph.add_run(
                source_run.text
            )

            _apply_formatting(
                target_run,
                source_run.formatting,
            )

    output.save(output_path)
